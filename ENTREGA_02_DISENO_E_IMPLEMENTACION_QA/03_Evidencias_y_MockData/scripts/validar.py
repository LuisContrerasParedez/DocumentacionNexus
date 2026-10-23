from pathlib import Path
from collections import Counter
import json,itertools,hashlib,zipfile,ast,re
from openpyxl import load_workbook
from docx import Document
B=Path(__file__).resolve().parents[1];E=B.parent;C=json.loads((B/'Casos_Diseno.json').read_text());M=json.loads((B/'MockData_Sintetica.json').read_text());PW=json.loads((B/'Pairwise_Cobertura.json').read_text());checks=[]
def ck(n,b):checks.append(dict(control=n,resultado='OK' if b else 'ERROR'))
ck('IDs de caso únicos',len({x['id'] for x in C})==len(C));ck('24 requerimientos presentes',{x['req'] for x in C}=={f'REQ-{n:02}' for n in range(1,25)})
ck('Condiciones y casos pertenecen a la misma REQ',all(x['condition'].split('-')[1]==x['id'].split('-')[1] and x['req'].replace('-','')==x['id'].split('-')[1] for x in C))
ck('Datos por caso únicos',len({x['data_id'] for x in C})==len(C));ck('Campos mínimos y pasos presentes',all(all(x[k] for k in ['id','req','title','technique','priority','preconditions','steps','data','expected']) for x in C));ck('Prioridad normalizada',all(x['priority'] in ['Alta','Media','Baja'] for x in C))
ck('Estados bloqueados coherentes',all(bool(x['blocker'])==(x['status']=='Bloqueado por definición') for x in C))
ck('Particiones, BVA, decisiones, estados y Pairwise presentes',all(any(t in x['technique'] for x in C) for t in ['Partición de equivalencia','Valores límite','Tabla de decisión','Transición de estados','Pairwise']))
factors=PW['factores'];names=list(factors);target={(names[i],a,names[j],b) for i in range(4) for j in range(i+1,4) for a in factors[names[i]] for b in factors[names[j]]};seen={(names[i],r[names[i]],names[j],r[names[j]]) for r in PW['filas'] for i in range(4) for j in range(i+1,4)};ck('Cobertura Pairwise independiente 24/24',seen==target and len(target)==24)
clientids={x['id'] for x in M['clientes']};acctids={x['id'] for x in M['cuentas']};ck('FK cliente→cuenta válidas',all(x['cliente_id'] in clientids for x in M['cuentas']));ck('FK cuenta→movimiento válidas',all(x['cuenta_id'] in acctids for x in M['movimientos']));ck('Volumen 8/1000/2000/10000',[len(M[k]) for k in ['usuarios','clientes','cuentas','movimientos']]==[8,1000,2000,10000]);ck('Emails reservados de prueba',all(x['email'].endswith('@example.invalid') for k in ['clientes','usuarios'] for x in M[k]));ck('Referencias/cuentas únicas',len({x['numero_cuenta'] for x in M['cuentas']})==2000 and len({x['referencia'] for x in M['movimientos']})==10000)
collection=json.loads((B/'Postman_Collection_QA.json').read_text());posts=[i for group in collection['item'] for i in group['item']];known={x['id'] for x in C}
ck('Colección JSON v2.1',collection['info']['schema'].endswith('/v2.1.0/collection.json'))
ck('Peticiones trazables',all(i['name'].split(' · ')[0] in known for i in posts if i['name'].startswith('TC-')))
ck('Guard de URL QA presente',any('skipRequest' in x for x in collection['event'][0]['script']['exec']))
for i in posts:
 if 'body' in i['request']:json.loads(i['request']['body']['raw'])
ck('Cuerpos JSON parseables',True)
rd=json.loads((B/'Capturas_Readiness_Entorno/Readiness_Resultados.json').read_text());ck('15 evidencias reales de readiness',len(rd['checks'])==15 and all(x['estado']=='APTO' for x in rd['checks']));ck('Huella de colección coincide con evidencia',next(x for x in rd['checks'] if x['id']=='RD-14')['observado']['sha256']==hashlib.sha256((B/'Postman_Collection_QA.json').read_bytes()).hexdigest())
wb=load_workbook(E/'02_Matrices_y_Casos_Prueba/E2_Matrices_y_CasosPrueba_EquipoXX.xlsx',data_only=False);ck('Matrices completas',all(wb[n].max_row==len(C)+1 for n in ['RTM','Casos_Prueba','Data_Matrix']))
doc=Document(E/'01_Documento_Consolidado_QA/E2_Documento_Diseno_Implementacion_EquipoXX.docx');text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]);ck('Todos los casos incluidos en Word',all(x['id'] in text for x in C));ck('Todos los datos incluidos en Word',all(x['data_id'] in text for x in C));ck('Sin JWT expuestos en reportes',not any(re.search(r'eyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}',p.read_text()) for p in (B/'Capturas_Readiness_Entorno').glob('*.json')))
for p in (B/'scripts').glob('*.py'):ast.parse(p.read_text())
ck('Scripts Python válidos',True)
events=json.loads((B/'Catalogo_Eventos_QA.json').read_text())['eventos']
ck('20 eventos únicos con casos y puntos de emisión',len(events)==20 and len({x['codigo'] for x in events})==20 and all(x['caso'] in known and x['puntos_emision'] for x in events))
ck('Sin casos bloqueados por definición',not any(x['blocker'] for x in C))
mockchecks=json.loads((B/'Capturas_Readiness_Entorno/Mock_y_Compuerta_Resultados.json').read_text())
ck('12 comprobaciones reales de mock y compuerta OK',len(mockchecks['controles'])==12 and mockchecks['errores']==0)
pdfpath=E/'01_Documento_Consolidado_QA/E2_Documento_Diseno_Implementacion_EquipoXX.pdf'
if pdfpath.exists():
 import pymupdf
 pdf=pymupdf.open(pdfpath);pdftext='\n'.join(p.get_text() for p in pdf)
 ck('PDF contiene todos los IDs de caso',all(x['id'] in pdftext for x in C))
 ck('PDF sin páginas vacías',all(len(p.get_text().strip())>100 for p in pdf))
 ck('PDF sin texto fuera de margen horizontal',not any(b[0]<35 or b[2]>p.rect.width-35 for p in pdf for b in p.get_text('blocks')))
report=dict(fecha='2026-10-06',casos=len(C),requerimientos=24,condiciones=len({x['condition'] for x in C}),bloqueados=sum(bool(x['blocker']) for x in C),pairwise_pares=len(seen),peticiones=len(posts),controles=checks,errores=sum(x['resultado']=='ERROR' for x in checks));(B/'Validacion_Artefactos.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in report.items() if k!='controles'},ensure_ascii=False,indent=2));raise SystemExit(bool(report['errores']))
